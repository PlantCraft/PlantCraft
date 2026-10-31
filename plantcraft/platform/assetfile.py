"""The asset file: the asset tree alone (no sources, no connections), in and out.

Assets › Import and Assets › Export move the tree between installations and other tools:
each asset with its type, parent, record (attribute values) and the signals it references by
role, and the relations between assets. Data sources and their connections stay in Settings
(the full configuration file, ``plantmap.py``).

Every format is read into one **document** and every export is written from it::

    {"format": "processcraft-assets", "version": 1,
     "assets": [{"asset_id": "FIC-1003", "name": "Recycle flow", "type": "loop",
                 "parent_id": "u-200-recycle", "description": "...",
                 "attributes": {"criticality": "high"},
                 "signals": [{"role": "pv", "tag": "CPM.Loops.FIC-1003.PV", "source_id": "plantcraft"}]}],
     "relations": [{"from": "FIC-1003", "type": "controls", "to": "FV-1003"}]}

Formats: that JSON; Excel (sheets ``assets``, ``signals``, ``relations``); CSV (the assets
sheet, a ``signal.<role>`` column per role); ISA-95 B2MML ``EquipmentInformation`` (MESA's
public schema: nested ``Equipment`` with ``EquipmentLevel``, ``EquipmentClassID`` and
``EquipmentProperty``); and the AF XML of PI Asset Framework (from its public documentation:
nested ``AFElement`` with ``Template`` and ``AFAttribute``, a ``PI Point`` data reference
naming the tag).

Importing is two steps, so a person checks the file before anything changes: ``review``
reads a document and says, asset by asset, what would be created or changed and what is wrong;
the page lets them fix names, types and parents, and ``apply`` writes the document they send
back. An import adds and updates; it never removes.
"""

from __future__ import annotations

import csv
import io
import json
import math
import re
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from typing import Any

from plantcraft.i18n import t
from plantcraft_sdk.assets import RELATION_TYPES

FORMAT = "processcraft-assets"
VERSION = 1
MAX_ASSETS = 20000
MAX_BYTES = 20 * 1024 * 1024
BASE_COLUMNS = ("asset_id", "name", "type", "parent_id", "description")
# Words a file may use for the type column.
TYPE_COLUMNS = ("type", "kind", "tipo")
SIGNAL_PREFIX = "signal."
B2MML_NS = "http://www.mesa.org/xml/B2MML"
# Export: the ISA-95 level of each root type; import: the type a level becomes when the
# file's EquipmentClassID names none of ours.
B2MML_LEVELS = {
    "site": "Site",
    "unit": "Area",
    "area": "ProcessCell",
    "equipment": "Unit",
    "instrument": "ControlModule",
    "loop": "ControlModule",
}
LEVEL_TYPES = {
    "area": "unit",
    "processcell": "area",
    "productionline": "area",
    "productionunit": "area",
    "storagezone": "area",
    "workcenter": "area",
    "unit": "equipment",
    "workunit": "equipment",
    "workcell": "equipment",
    "storageunit": "equipment",
    "equipmentmodule": "equipment",
    "controlmodule": "instrument",
    "other": "equipment",
}
# Levels above the plant: their children hang from this installation's plant.
ROOT_LEVELS = {"enterprise", "site"}
FORMATS = {
    "json": ("application/json", "json"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "csv": ("text/csv; charset=utf-8", "csv"),
    "b2mml": ("application/xml", "b2mml.xml"),
    "af": ("application/xml", "af.xml"),
}


class AssetFileError(ValueError):
    """The file cannot be read at all."""


# Reading values ------------------------------------------------------------------


def _text(value: Any) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.strip().lower()).strip("_")


def _blank_document() -> dict[str, Any]:
    return {"format": FORMAT, "version": VERSION, "assets": [], "relations": []}


def _asset(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "asset_id": _text(row.get("asset_id")),
        "name": _text(row.get("name")),
        "type": _text(next((row[key] for key in TYPE_COLUMNS if _text(row.get(key))), "")),
        "parent_id": _text(row.get("parent_id")),
        "description": _text(row.get("description")),
        "attributes": dict(row.get("attributes") or {}),
        "signals": list(row.get("signals") or []),
    }


# Reading files --------------------------------------------------------------------


def read(content: bytes, filename: str) -> tuple[dict[str, Any], list[str]]:
    """The document a file holds, and notes on what was read (the format, what was skipped)."""

    if len(content) > MAX_BYTES:
        raise AssetFileError(t("The file is larger than {size} MB.", size=MAX_BYTES // 2**20))
    name = (filename or "").lower()
    if name.endswith(".json"):
        return _read_json(content)
    if name.endswith(".xlsx"):
        return _read_excel(content)
    if name.endswith(".csv"):
        return _read_csv(content)
    if name.endswith(".xml"):
        return _read_xml(content)
    raise AssetFileError(t("Use a .json, .xlsx, .csv or .xml file."))


def _read_json(content: bytes) -> tuple[dict[str, Any], list[str]]:
    try:
        data = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise AssetFileError(t("The file is not valid JSON: {error}", error=exc)) from exc
    if isinstance(data, list):
        data = {"assets": data}
    if not isinstance(data, dict):
        raise AssetFileError(t("The JSON file must hold an object with an assets list."))
    if data.get("format") == "processcraft-map" or "sources" in data or "loops" in data:
        raise AssetFileError(
            t(
                "This is a full configuration file (sources, signals and assets): import it in "
                "Settings, Configuration file."
            )
        )
    assets = data.get("assets")
    if not isinstance(assets, list):
        raise AssetFileError(t("The JSON file must hold an object with an assets list."))
    document = _blank_document()
    for row in assets:
        if isinstance(row, dict):
            item = _asset(row)
            item["signals"] = [
                {
                    "role": _text(signal.get("role")),
                    "tag": _text(signal.get("tag") or signal.get("signal_id")),
                    "source_id": _text(signal.get("source_id")),
                }
                for signal in row.get("signals") or []
                if isinstance(signal, dict)
            ]
            document["assets"].append(item)
    document["relations"] = [
        {
            "from": _text(row.get("from") or row.get("from_id")),
            "type": _text(row.get("type") or row.get("relation")),
            "to": _text(row.get("to") or row.get("to_id")),
        }
        for row in data.get("relations") or []
        if isinstance(row, dict)
    ]
    return document, [t("PlantCraft assets file (JSON).")]


def _row_asset(row: dict[str, Any], *, default_source: str = "") -> dict[str, Any]:
    """A spreadsheet row: the base columns, ``signal.<role>`` columns, the rest attributes."""

    item = _asset(row)
    source = _text(row.get("source_id")) or default_source
    for key, value in row.items():
        key = str(key).strip()
        if not key or key == "_line" or key in BASE_COLUMNS or key in TYPE_COLUMNS:
            continue
        if key in {"source_id", "signals", "attributes"}:
            continue
        if key.lower().startswith(SIGNAL_PREFIX):
            if _text(value):
                item["signals"].append(
                    {"role": key[len(SIGNAL_PREFIX) :], "tag": _text(value), "source_id": source}
                )
            continue
        if _text(value) != "":
            item["attributes"][key] = value if isinstance(value, (int, float)) else _text(value)
    return item


def _read_csv(content: bytes) -> tuple[dict[str, Any], list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
    if not rows or "asset_id" not in {str(key).strip().lower() for key in rows[0]}:
        raise AssetFileError(t("The CSV file needs an asset_id column."))
    document = _blank_document()
    for row in rows:
        clean = {str(key).strip(): value for key, value in row.items() if key is not None}
        clean = {
            (key.lower() if key.lower() in BASE_COLUMNS else key): v for key, v in clean.items()
        }
        if _text(clean.get("asset_id")):
            document["assets"].append(_row_asset(clean))
    return document, [t("CSV: one row per asset.")]


def _read_excel(content: bytes) -> tuple[dict[str, Any], list[str]]:
    try:
        import openpyxl
    except ImportError as exc:  # the server extra installs it
        raise AssetFileError(t("Excel files need openpyxl on the server.")) from exc
    try:
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise AssetFileError(t("The Excel file could not be opened: {error}", error=exc)) from exc
    sheets: dict[str, list[dict[str, Any]]] = {}
    for sheet in book.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        keys = [_text(cell) for cell in header]
        found = []
        for values in rows:
            row = {key: value for key, value in zip(keys, values, strict=False) if key}
            if any(_text(value) for value in row.values()):
                found.append(row)
        sheets[sheet.title.strip().lower()] = found
    if "assets" not in sheets:
        raise AssetFileError(t("The workbook needs a sheet named assets."))
    document = _blank_document()
    by_id: dict[str, dict[str, Any]] = {}
    for row in sheets["assets"]:
        clean = {
            (str(key).lower() if str(key).lower() in BASE_COLUMNS else key): v
            for key, v in row.items()
        }
        if _text(clean.get("asset_id")):
            item = _row_asset(clean)
            document["assets"].append(item)
            by_id.setdefault(item["asset_id"], item)
    for row in sheets.get("signals", []):
        owner = by_id.get(_text(row.get("asset_id")))
        if owner is not None and _text(row.get("tag")):
            owner["signals"].append(
                {
                    "role": _text(row.get("role")),
                    "tag": _text(row.get("tag")),
                    "source_id": _text(row.get("source_id")),
                }
            )
    document["relations"] = [
        {
            "from": _text(row.get("from_id") or row.get("from")),
            "type": _text(row.get("relation") or row.get("type")),
            "to": _text(row.get("to_id") or row.get("to")),
        }
        for row in sheets.get("relations", [])
    ]
    return document, [t("Excel workbook: sheets assets, signals and relations.")]


def _local(tag: str) -> str:
    """An element's name without its namespace."""

    return tag.rsplit("}", 1)[-1]


def _child(element: ET.Element, name: str) -> ET.Element | None:
    return next((item for item in element if _local(item.tag) == name), None)


def _children(element: ET.Element, name: str) -> list[ET.Element]:
    return [item for item in element if _local(item.tag) == name]


def _child_text(element: ET.Element, name: str) -> str:
    found = _child(element, name)
    return (found.text or "").strip() if found is not None else ""


def _read_xml(content: bytes) -> tuple[dict[str, Any], list[str]]:
    head = content[:4096].decode("utf-8", "ignore").upper()
    if "<!DOCTYPE" in head or "<!ENTITY" in head:
        raise AssetFileError(t("XML files with a DOCTYPE or entities are not read."))
    try:
        root = ET.fromstring(content)
    except ET.ParseError as exc:
        raise AssetFileError(t("The XML file could not be read: {error}", error=exc)) from exc
    name = _local(root.tag)
    if (
        name == "AF"
        or name == "AFDatabase"
        or any(_local(item.tag) == "AFElement" for item in root.iter())
    ):
        return _read_af(root)
    if name in {
        "EquipmentInformation",
        "Equipment",
        "ShowEquipmentInformation",
        "SyncEquipmentInformation",
    } or any(_local(item.tag) == "Equipment" for item in root.iter()):
        return _read_b2mml(root)
    raise AssetFileError(
        t("This XML is neither ISA-95 B2MML equipment nor PI AF XML (no Equipment or AFElement).")
    )


def _unique(asset_id: str, parent: str, seen: set[str], notes: list[str]) -> str:
    """Element names repeat across branches in other tools: the second one gets its parent's."""

    if asset_id not in seen:
        seen.add(asset_id)
        return asset_id
    candidate = f"{parent}.{asset_id}" if parent else asset_id
    index = 2
    while candidate in seen:
        candidate = f"{parent}.{asset_id}.{index}"
        index += 1
    seen.add(candidate)
    notes.append(t("{name} appears twice; the second one is {id}.", name=asset_id, id=candidate))
    return candidate


def _read_b2mml(root: ET.Element) -> tuple[dict[str, Any], list[str]]:
    document = _blank_document()
    notes = [t("ISA-95 B2MML equipment (EquipmentInformation).")]
    seen: set[str] = set()
    top = [root] if _local(root.tag) == "Equipment" else _children(root, "Equipment")
    if not top:
        top = [item for item in root.iter() if _local(item.tag) == "Equipment"][:1]

    def walk(element: ET.Element, parent: str) -> None:
        raw_id = _child_text(element, "ID")
        level = _child_text(element, "EquipmentLevel")
        if level.strip().lower() in ROOT_LEVELS:
            for child in _children(element, "Equipment"):
                walk(child, parent)
            notes.append(
                t(
                    "{id} ({level}) is this installation's plant: what is in it goes under the plant.",
                    id=raw_id or "?",
                    level=level,
                )
            )
            return
        if not raw_id:
            return
        asset_id = _unique(raw_id.replace(" ", "-"), parent, seen, notes)
        attributes: dict[str, Any] = {}
        signals: list[dict[str, str]] = []
        description = ""
        for prop in _children(element, "EquipmentProperty"):
            key = _child_text(prop, "ID")
            value_element = _child(prop, "Value")
            value = _child_text(value_element, "ValueString") if value_element is not None else ""
            if not value:
                value = _child_text(prop, "Value")
            if not key or value == "":
                continue
            if key.startswith(SIGNAL_PREFIX):
                source, _, tag = value.partition("|")
                signals.append(
                    {
                        "role": key[len(SIGNAL_PREFIX) :],
                        "tag": tag or source,
                        "source_id": source if tag else "",
                    }
                )
            elif key == "processcraft.description":
                description = value
            else:
                attributes[key] = value
        name = _child_text(element, "Description") or raw_id
        document["assets"].append(
            {
                "asset_id": asset_id,
                "name": name,
                "type": _child_text(element, "EquipmentClassID")
                or LEVEL_TYPES.get(level.strip().lower(), "equipment"),
                "parent_id": parent,
                "description": description,
                "attributes": attributes,
                "signals": signals,
                "level": level,
            }
        )
        for child in _children(element, "Equipment"):
            walk(child, asset_id)

    for element in top:
        walk(element, "")
    return document, notes


def _af_tag(config: str, element: str) -> str:
    """The tag a PI Point data reference names: ``\\\\server\\tag;options`` → ``tag``."""

    path = config.split(";", 1)[0].split("?", 1)[0].strip()
    tag = path.rsplit("\\", 1)[-1].strip()
    return tag.replace("%Element%", element).replace("%element%", element)


def _read_af(root: ET.Element) -> tuple[dict[str, Any], list[str]]:
    document = _blank_document()
    notes = [t("PI AF XML (elements, templates and attributes).")]
    seen: set[str] = set()
    unresolved: list[str] = []

    def walk(element: ET.Element, parent: str) -> None:
        raw = _child_text(element, "Name")
        if not raw:
            return
        asset_id = _unique(raw.replace(" ", "-"), parent, seen, notes)
        attributes: dict[str, Any] = {}
        signals: list[dict[str, str]] = []
        description = ""
        for attribute in _children(element, "AFAttribute"):
            key = _child_text(attribute, "Name")
            if not key:
                continue
            reference = _child_text(attribute, "DataReference")
            if reference:
                config = _child_text(attribute, "ConfigString")
                if reference.lower() == "pi point" and config:
                    tag = _af_tag(config, raw)
                    if "%" in tag:
                        unresolved.append(f"{raw}.{key}")
                        continue
                    server = config.split(";", 1)[0].strip().lstrip("\\").split("\\", 1)[0]
                    signals.append(
                        {
                            "role": _slug(key),
                            "tag": tag,
                            "source_id": "" if "%" in server else server,
                        }
                    )
                continue
            value = _child_text(attribute, "Value")
            if key.lower() in {"notes", "description"}:
                description = value
            elif value != "":
                attributes[key] = value
        document["assets"].append(
            {
                "asset_id": asset_id,
                "name": _child_text(element, "Description") or raw,
                "type": _child_text(element, "Template") or "equipment",
                "parent_id": parent,
                "description": description,
                "attributes": attributes,
                "signals": signals,
            }
        )
        for child in _children(element, "AFElement"):
            walk(child, asset_id)

    databases = (
        [root] if _local(root.tag) == "AFDatabase" else _children(root, "AFDatabase") or [root]
    )
    for database in databases:
        for element in _children(database, "AFElement"):
            walk(element, "")
    if unresolved:
        notes.append(
            t(
                "Data references with substitutions that cannot be resolved were skipped: {items}.",
                items=", ".join(unresolved[:10]) + (" …" if len(unresolved) > 10 else ""),
            )
        )
    return document, notes


# Reviewing ------------------------------------------------------------------------


def _type_of(types: Any, value: str) -> str | None:
    """A type id from what a file says: the id, or the name in any case."""

    if not value:
        return None
    if types.exists(value):
        return value
    wanted = value.strip().lower()
    for item in types.definitions().values():
        if item["id"].lower() == wanted or str(item["name"]).strip().lower() == wanted:
            return item["id"]
    slug = _slug(value)
    return slug if slug and types.exists(slug) else None


def _attribute_entry(resolved: dict[str, Any], key: str) -> dict[str, Any] | None:
    wanted = str(key).strip().lower()
    for entry in resolved.get("attributes", []):
        if entry["key"] == wanted or str(entry["name"]).strip().lower() == wanted:
            return entry
    return None


def _role_key(resolved: dict[str, Any] | None, role: str) -> str:
    wanted = str(role).strip().lower()
    for entry in (resolved or {}).get("roles", []):
        if entry["key"] == wanted or str(entry["name"]).strip().lower() == wanted:
            return entry["key"]
    return _slug(role) if role else ""


def review(platform: Any, document: dict[str, Any], *, anchor: str | None = None) -> dict[str, Any]:
    """What applying ``document`` would do, asset by asset, and what is wrong with it.

    Each asset comes back with its ``status`` (create, update, unchanged or error), the
    ``changes`` it makes (fields, attribute keys, signals), its ``errors`` and ``warnings``,
    and the values as they would be written (type resolved, attributes checked, parent
    placed). ``anchor`` is where assets without a known parent go: the plant by default.
    """

    from plantcraft.assets.types import AssetTypeError, check_value

    store = platform.workspace.assets
    types = store.types
    root = platform.root_asset_id
    anchor = anchor if anchor and store.get(anchor) is not None else root
    issues: list[dict[str, Any]] = []
    raw_assets = [item for item in document.get("assets") or [] if isinstance(item, dict)]
    if len(raw_assets) > MAX_ASSETS:
        raise AssetFileError(t("The file has more than {count} assets.", count=MAX_ASSETS))
    in_file: dict[str, dict[str, Any]] = {}
    rows: list[dict[str, Any]] = []
    known_tags = _known_tags(platform)
    known_sources = _known_sources(platform)
    for index, raw in enumerate(raw_assets):
        item = _asset(raw)
        item["signals"] = [
            {
                "role": _text(signal.get("role")),
                "tag": _text(signal.get("tag")),
                "source_id": _text(signal.get("source_id")),
            }
            for signal in raw.get("signals") or []
            if isinstance(signal, dict) and _text(signal.get("tag"))
        ]
        item["include"] = raw.get("include", True) is not False
        item["index"] = index
        if isinstance(raw.get("draft"), dict):
            # A drafted plant's row (plantdraft.py): its confidence and sources, carried as is.
            item["draft"] = raw["draft"]
        errors: list[str] = []
        warnings: list[str] = []
        asset_id = item["asset_id"]
        if not asset_id:
            errors.append(t("Enter the asset id."))
        elif any(char.isspace() for char in asset_id) or len(asset_id) > 128:
            errors.append(t("The asset_id '{asset}' cannot have spaces.", asset=asset_id[:60]))
        elif asset_id == root:
            errors.append(t("{asset} is the plant, the root of the tree.", asset=asset_id))
        elif asset_id in in_file:
            errors.append(t("Asset {asset} appears more than once.", asset=asset_id))
        item["name"] = item["name"] or asset_id
        type_id = _type_of(types, item["type"])
        if type_id is None:
            errors.append(
                t(
                    "Unknown asset type: {type}. Choose one of this installation's types.",
                    type=item["type"] or "—",
                )
            )
        elif type_id == "site":
            errors.append(t("The plant is the root of the tree; choose another type."))
        else:
            item["type"] = type_id
        resolved = types.resolved(type_id) if type_id else None
        values: dict[str, Any] = {}
        ignored: list[str] = []
        for key, raw_value in (item["attributes"] or {}).items():
            entry = _attribute_entry(resolved or {}, key)
            if entry is None:
                ignored.append(str(key)[:40])
                continue
            try:
                value = check_value(entry, raw_value)
            except AssetTypeError as exc:
                errors.append(str(exc))
                continue
            if value is not None:
                values[entry["key"]] = value
        item["attributes"] = values
        if ignored and resolved:
            warnings.append(
                t(
                    "Not attributes of {type}, left out: {keys}.",
                    type=resolved["name"],
                    keys=", ".join(sorted(ignored)[:12]),
                )
            )
        signals = []
        for signal in item["signals"]:
            role = _role_key(resolved, signal["role"])
            if signal["source_id"] and signal["source_id"] not in known_sources:
                warnings.append(
                    t(
                        "Source {source} is not configured here; {tag} is read from the source that has it.",
                        source=signal["source_id"][:60],
                        tag=signal["tag"][:120],
                    )
                )
                signal["source_id"] = ""
            source = signal["source_id"] or (
                platform.sources.source_of(signal["tag"]) if signal["tag"] in known_tags else ""
            )
            if signal["tag"] not in known_tags:
                warnings.append(
                    t(
                        "Signal {tag} is not in the platform yet; it is referenced anyway.",
                        tag=signal["tag"][:120],
                    )
                )
            if not any(
                item2["tag"] == signal["tag"] and item2["role"] == role for item2 in signals
            ):
                signals.append({"role": role, "tag": signal["tag"], "source_id": source})
        item["signals"] = signals
        item["errors"] = errors
        item["warnings"] = warnings
        rows.append(item)
        if asset_id and asset_id not in in_file:
            in_file[asset_id] = item
    included = {row["asset_id"] for row in rows if row["include"] and row["asset_id"]}
    # Parents: in the file (and included), in the tree, else the anchor.
    for row in rows:
        parent = row["parent_id"]
        if not parent or parent == root:
            row["parent_id"] = anchor if not parent else root
            continue
        if parent in included:
            continue
        if store.get(parent) is None:
            row["warnings"].append(
                t(
                    "Parent {parent} is not in the file or the tree; it goes under {anchor}.",
                    parent=parent[:120],
                    anchor=anchor,
                )
            )
            row["parent_id"] = anchor
    by_id = {row["asset_id"]: row for row in rows if row["include"] and row["asset_id"]}
    for row in rows:
        seen = {row["asset_id"]}
        parent = row["parent_id"]
        while parent in by_id:
            if parent in seen:
                row["errors"].append(
                    t(
                        "Asset {asset} ends up inside itself (cycle in parent_id).",
                        asset=row["asset_id"],
                    )
                )
                break
            seen.add(parent)
            parent = by_id[parent]["parent_id"]
    # Asset references in records may point at the file or the tree.
    for row in rows:
        resolved = types.resolved(row["type"]) if types.exists(row["type"]) else None
        for entry in (resolved or {}).get("attributes", []):
            value = row["attributes"].get(entry["key"])
            if (
                entry["data_type"] == "asset"
                and value
                and value not in by_id
                and store.get(value) is None
            ):
                row["errors"].append(
                    t("{key}: asset {asset} not found.", key=entry["key"], asset=str(value)[:120])
                )
    # What changes.
    for row in rows:
        if row["errors"]:
            row["status"] = "error"
            row["changes"] = []
            continue
        current = store.get(row["asset_id"])
        if current is None:
            row["status"] = "create"
            row["changes"] = []
            continue
        changes = [
            field
            for field, wanted in (
                ("name", row["name"]),
                ("type", row["type"]),
                ("parent_id", row["parent_id"]),
                ("description", row["description"]),
            )
            if current.get("kind" if field == "type" else field) != wanted
        ]
        if field_origin := current.get("origin"):
            if field_origin not in {"manual", "map"} and changes:
                row["warnings"].append(
                    t(
                        "{asset} came from {origin}; after the import it is kept by hand.",
                        asset=row["asset_id"],
                        origin=field_origin,
                    )
                )
        stored = store.stored_attributes([row["asset_id"]]).get(row["asset_id"], {})
        changes += [
            f"attributes.{key}"
            for key, value in row["attributes"].items()
            if stored.get(key) != value
        ]
        bound = {(item["signal_id"], item["role"]) for item in store.signals(row["asset_id"])}
        changes += [
            f"signals.{item['role'] or item['tag']}"
            for item in row["signals"]
            if (item["tag"], item["role"]) not in bound
        ]
        row["changes"] = changes
        row["status"] = "update" if changes else "unchanged"
    relations = []
    excluded = {row["asset_id"] for row in rows if not row["include"]}
    for raw in document.get("relations") or []:
        if not isinstance(raw, dict):
            continue
        relation = {
            "from": _text(raw.get("from")),
            "type": _text(raw.get("type")),
            "to": _text(raw.get("to")),
            "include": raw.get("include", True) is not False,
        }
        problems = []
        left_out = [
            end
            for end in (relation["from"], relation["to"])
            if end in excluded and store.get(end) is None
        ]
        if left_out:
            # An asset left out of the import takes its relations with it.
            relation.update(include=False, errors=[], status="skipped")
            relations.append(relation)
            continue
        if relation["type"] not in RELATION_TYPES:
            problems.append(t("Unknown relation: use {types}.", types=", ".join(RELATION_TYPES)))
        for end in (relation["from"], relation["to"]):
            if end not in by_id and store.get(end) is None:
                problems.append(t("Asset {asset} not found.", asset=end[:120] or "—"))
        if relation["from"] == relation["to"]:
            problems.append(t("An asset cannot be related to itself."))
        current = (
            {
                (item["from_id"], item["relation"], item["to_id"])
                for item in store.relations(relation["from"])
            }
            if not problems
            else set()
        )
        relation["errors"] = problems
        relation["status"] = (
            "error"
            if problems
            else (
                "unchanged"
                if (relation["from"], relation["type"], relation["to"]) in current
                else "create"
            )
        )
        relations.append(relation)
    counts = {status: 0 for status in ("create", "update", "unchanged", "error")}
    for row in rows:
        if row["include"]:
            counts[row["status"]] += 1
    relation_counts = {status: 0 for status in ("create", "unchanged", "error")}
    for relation in relations:
        if relation["include"]:
            relation_counts[relation["status"]] += 1
    errors = counts["error"] + relation_counts["error"]
    return {
        "format": FORMAT,
        "version": VERSION,
        "anchor": anchor,
        "assets": rows,
        "relations": relations,
        "counts": counts,
        "relation_counts": relation_counts,
        "issues": issues,
        "can_apply": errors == 0
        and (counts["create"] + counts["update"] + relation_counts["create"]) > 0,
    }


def _known_tags(platform: Any) -> set[str]:
    sources = platform.sources
    found = set(sources.routes()) | set(platform.scenario.source.tags)
    found |= set(platform.workspace.assets.bound_tags())
    return found


def _known_sources(platform: Any) -> set[str]:
    found = {item["source_id"] for item in platform.sources.list()}
    layer = getattr(platform, "datalayer", None)
    if layer is not None and getattr(layer, "manages_sources", False):
        try:
            found |= {item["source_id"] for item in layer.sources()}
        except Exception:  # noqa: BLE001 - Historian down: the local sources count
            pass
    return found


def apply(
    platform: Any, document: dict[str, Any], *, actor: str, anchor: str | None = None
) -> dict[str, Any]:
    """Write a reviewed document (the included rows); nothing is written when any has an error."""

    result = review(platform, document, anchor=anchor)
    if not result["can_apply"]:
        return {**result, "applied": False}
    store = platform.workspace.assets
    rows = [
        row for row in result["assets"] if row["include"] and row["status"] in {"create", "update"}
    ]
    by_id = {row["asset_id"]: row for row in rows}
    ordered: list[dict[str, Any]] = []
    placed: set[str] = set()

    def place(row: dict[str, Any]) -> None:
        if row["asset_id"] in placed:
            return
        placed.add(row["asset_id"])
        parent = by_id.get(row["parent_id"])
        if parent is not None:
            place(parent)
        ordered.append(row)

    for row in rows:
        place(row)
    # What this import changes, kept so it can be undone (``undo``): every asset it creates, and
    # the name, place and values of every asset it overwrites, as they are now.
    undo: dict[str, Any] = {"created": [], "updated": {}, "signals": [], "relations": []}
    before = store.stored_attributes([row["asset_id"] for row in ordered])
    for row in ordered:
        current = store.get(row["asset_id"])
        if current is None:
            undo["created"].append(row["asset_id"])
        else:
            undo["updated"][row["asset_id"]] = {
                "name": current["name"],
                "kind": current["kind"],
                "parent_id": current["parent_id"],
                "description": current["description"],
                "position": current["position"],
                "attributes": {
                    key: before.get(row["asset_id"], {}).get(key) for key in row["attributes"]
                },
            }
    for row in ordered:
        current = store.get(row["asset_id"])
        store.save(
            row["asset_id"],
            name=row["name"],
            kind=row["type"],
            parent_id=row["parent_id"],
            description=row["description"],
            origin="manual",
            position=(current or {}).get("position", row["index"]),
        )
        stored = store.stored_attributes([row["asset_id"]]).get(row["asset_id"], {})
        values = {
            key: value for key, value in row["attributes"].items() if stored.get(key) != value
        }
        if values:
            changes = store.set_attributes(row["asset_id"], values, actor=actor, check_refs=False)
            platform.record_attribute_changes(row["asset_id"], actor, changes)
        bound = {(item["signal_id"], item["role"]) for item in store.signals(row["asset_id"])}
        new = [item for item in row["signals"] if (item["tag"], item["role"]) not in bound]
        if new:
            undo["signals"] += [[row["asset_id"], item["tag"]] for item in new]
            store.bind_signals(
                row["asset_id"],
                [
                    {"signal_id": item["tag"], "role": item["role"], "source_id": item["source_id"]}
                    for item in new
                ],
            )
    for relation in result["relations"]:
        if relation["include"] and relation["status"] == "create":
            store.relate(relation["from"], relation["type"], relation["to"], actor=actor)
            undo["relations"].append([relation["from"], relation["type"], relation["to"]])
    summary = {
        "created": result["counts"]["create"],
        "updated": result["counts"]["update"],
        "relations": result["relation_counts"]["create"],
    }
    import_id = platform.workspace.asset_imports.record(actor, summary, undo)
    platform.workspace.audit.record(
        actor, "assets.imported", "assets", {**summary, "import_id": import_id}
    )
    done = review(platform, document, anchor=anchor)
    return {
        **done,
        "applied": True,
        "created": result["counts"]["create"],
        "updated": result["counts"]["update"],
        "related": result["relation_counts"]["create"],
    }


def start_scope(platform: Any) -> dict[str, Any] | None:
    """What the newest import wrote, as the apps' first steps read it (``[[contributes.start]]``,
    ``plantcraft_sdk.contributes.StartIn``): its assets as they are now, their signals, and the
    units to work by (its assets of kind ``unit``, else its top assets). ``None`` when there is no
    import, or it was undone."""

    found = platform.workspace.asset_imports.last()
    if found is None or found["undone_at"] is not None:
        return None
    undo, store = found["undo"], platform.workspace.assets
    ids = [*undo.get("created", []), *undo.get("updated", {})]
    assets = [item for item in (store.get(asset_id) for asset_id in dict.fromkeys(ids)) if item]
    kept = {item["asset_id"] for item in assets}
    signals = [
        {
            "asset_id": item["asset_id"],
            "tag": item["signal_id"],
            "role": item.get("role") or "",
            "source_id": item.get("source_id") or "",
        }
        for item in store.bindings(origin="manual")
        if item["asset_id"] in kept
    ]
    units = [item["asset_id"] for item in assets if item["kind"] == "unit"] or [
        item["asset_id"] for item in assets if item.get("parent_id") not in kept
    ]
    return {
        "import_id": str(found["import_id"]),
        "assets": [
            {
                "id": item["asset_id"],
                "kind": item["kind"],
                "name": item["name"],
                "parent_id": item.get("parent_id"),
            }
            for item in assets
        ],
        "signals": signals,
        "units": units,
    }


# Undoing the last import ----------------------------------------------------------


def last_import(platform: Any) -> dict[str, Any] | None:
    """The newest import, whether it can still be undone, without its undo journal."""

    found = platform.workspace.asset_imports.last()
    if found is None:
        return None
    undo = found.pop("undo")
    return {**found, "can_undo": found["undone_at"] is None and bool(_journal_size(undo))}


def _journal_size(undo: dict[str, Any]) -> int:
    return sum(len(undo.get(key) or []) for key in ("created", "updated", "signals", "relations"))


def undo_last(platform: Any, *, actor: str) -> dict[str, Any]:
    """Take the last import back: the relations and signal bindings it added, the values it
    overwrote, then the assets it created (newest first; their children made by hand since then
    move up to the parent, as a delete does). Raises ``AssetFileError`` when there is none."""

    imports = platform.workspace.asset_imports
    found = imports.last()
    if found is None or found["undone_at"] is not None:
        raise AssetFileError(t("There is no import to undo."))
    undo, store = found["undo"], platform.workspace.assets
    for from_id, relation, to_id in undo.get("relations", []):
        store.unrelate(from_id, relation, to_id)
    if undo.get("signals"):
        store.unbind_signals([(asset_id, tag) for asset_id, tag in undo["signals"]])
    restored = 0
    for asset_id, previous in undo.get("updated", {}).items():
        if store.get(asset_id) is None:
            continue
        store.save(
            asset_id,
            name=previous["name"],
            kind=previous["kind"],
            parent_id=previous["parent_id"],
            description=previous["description"],
            origin="manual",
            position=previous["position"],
        )
        if previous["attributes"]:
            try:
                store.set_attributes(
                    asset_id, previous["attributes"], actor=actor, check_refs=False
                )
            except (KeyError, ValueError):
                pass  # A value the asset's type no longer takes stays as the import left it.
        restored += 1
    removed = 0
    for asset_id in reversed(undo.get("created", [])):
        if store.get(asset_id) is not None:
            store.delete(asset_id)
            removed += 1
    imports.mark_undone(found["import_id"], actor)
    detail = {"import_id": found["import_id"], "removed": removed, "restored": restored}
    platform.workspace.audit.record(actor, "assets.import_undone", "assets", detail)
    return detail


# Exporting ------------------------------------------------------------------------


def document(platform: Any, root: str | None = None) -> dict[str, Any]:
    """The tree (or the branch under ``root``, itself included) as a document."""

    store = platform.workspace.assets
    plant = platform.root_asset_id
    items = store.list()
    if root and root != plant and store.get(root) is not None:
        keep = store.descendants(root)
        items = [item for item in items if item["asset_id"] in keep]
    by_id = {item["asset_id"]: item for item in items}
    children: dict[str | None, list[dict[str, Any]]] = {}
    for item in items:
        if item["asset_id"] == plant:
            continue
        parent = (
            item["parent_id"] if item["parent_id"] in by_id and item["parent_id"] != plant else None
        )
        children.setdefault(parent, []).append(item)
    ordered: list[dict[str, Any]] = []

    def walk(parent: str | None) -> None:
        for item in sorted(
            children.get(parent, []), key=lambda row: (row["position"], row["name"])
        ):
            ordered.append(item)
            walk(item["asset_id"])

    walk(None)
    values = store.stored_attributes([item["asset_id"] for item in ordered])
    out = _blank_document()
    out["exported_at"] = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    for item in ordered:
        parent = (
            item["parent_id"] if item["parent_id"] in by_id and item["parent_id"] != plant else ""
        )
        out["assets"].append(
            {
                "asset_id": item["asset_id"],
                "name": item["name"],
                "type": item["kind"],
                "parent_id": parent,
                "description": item["description"],
                "attributes": values.get(item["asset_id"], {}),
                "signals": [
                    {
                        "role": signal["role"],
                        "tag": signal["signal_id"],
                        "source_id": signal.get("source_id") or "",
                    }
                    for signal in store.signals(item["asset_id"])
                ],
            }
        )
    exported = {item["asset_id"] for item in ordered}
    out["relations"] = [
        {"from": item["from_id"], "type": item["relation"], "to": item["to_id"]}
        for item in store.relations()
        if item["from_id"] in exported and item["to_id"] in exported
    ]
    return out


def export(platform: Any, kind: str, root: str | None = None) -> tuple[bytes, str, str]:
    """``(content, media type, file extension)`` of the tree in one format."""

    if kind not in FORMATS:
        raise AssetFileError(t("Unknown format."))
    data = document(platform, root)
    media, extension = FORMATS[kind]
    if kind == "json":
        content = json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")
    elif kind == "csv":
        content = _csv(data)
    elif kind == "xlsx":
        content = _xlsx(data)
    elif kind == "b2mml":
        content = _b2mml(platform, data)
    else:
        content = _af(platform, data)
    return content, media, extension


def _attribute_keys(data: dict[str, Any]) -> list[str]:
    return list(dict.fromkeys(key for item in data["assets"] for key in item["attributes"]))


def _cell(value: Any) -> Any:
    return "" if value is None else value


def _csv(data: dict[str, Any]) -> bytes:
    roles = list(
        dict.fromkeys(
            signal["role"] or "signal" for item in data["assets"] for signal in item["signals"]
        )
    )
    keys = _attribute_keys(data)
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(
        [*BASE_COLUMNS, "source_id", *(f"{SIGNAL_PREFIX}{role}" for role in roles), *keys]
    )
    for item in data["assets"]:
        by_role: dict[str, str] = {}
        for signal in item["signals"]:
            by_role.setdefault(signal["role"] or "signal", signal["tag"])
        sources = {signal["source_id"] for signal in item["signals"] if signal["source_id"]}
        writer.writerow(
            [
                *(item[key] for key in BASE_COLUMNS),
                next(iter(sources)) if len(sources) == 1 else "",
                *(by_role.get(role, "") for role in roles),
                *(_cell(item["attributes"].get(key)) for key in keys),
            ]
        )
    return ("﻿" + out.getvalue()).encode("utf-8")


def _xlsx(data: dict[str, Any]) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError as exc:
        raise AssetFileError(t("Excel files need openpyxl on the server.")) from exc
    book = openpyxl.Workbook()
    keys = _attribute_keys(data)
    sheet = book.active
    sheet.title = "assets"
    sheet.append([*BASE_COLUMNS, *keys])
    for item in data["assets"]:
        sheet.append(
            [
                *(item[key] for key in BASE_COLUMNS),
                *(_cell(item["attributes"].get(key)) for key in keys),
            ]
        )
    signals = book.create_sheet("signals")
    signals.append(["asset_id", "role", "tag", "source_id"])
    for item in data["assets"]:
        for signal in item["signals"]:
            signals.append([item["asset_id"], signal["role"], signal["tag"], signal["source_id"]])
    relations = book.create_sheet("relations")
    relations.append(["from_id", "relation", "to_id"])
    for relation in data["relations"]:
        relations.append([relation["from"], relation["type"], relation["to"]])
    for each in book.worksheets:
        for cell in each[1]:
            cell.font = Font(bold=True)
        each.freeze_panes = "A2"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _xml(root: ET.Element) -> bytes:
    ET.indent(root)
    return b'<?xml version="1.0" encoding="UTF-8"?>\n' + ET.tostring(root, encoding="utf-8")


def _tree(data: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    children: dict[str, list[dict[str, Any]]] = {}
    for item in data["assets"]:
        children.setdefault(item["parent_id"], []).append(item)
    return children


def _b2mml(platform: Any, data: dict[str, Any]) -> bytes:
    types = platform.workspace.assets.types
    ET.register_namespace("", B2MML_NS)
    q = lambda name: f"{{{B2MML_NS}}}{name}"  # noqa: E731
    root = ET.Element(q("EquipmentInformation"))
    ET.SubElement(root, q("ID")).text = FORMAT
    ET.SubElement(root, q("PublishedDate")).text = data.get("exported_at", "")
    children = _tree(data)
    site = platform.workspace.assets.get(platform.root_asset_id) or {}
    plant = ET.SubElement(root, q("Equipment"))
    ET.SubElement(plant, q("ID")).text = platform.root_asset_id
    ET.SubElement(plant, q("Description")).text = site.get("name", "")
    ET.SubElement(plant, q("EquipmentLevel")).text = "Site"

    def add(parent: ET.Element, item: dict[str, Any]) -> None:
        element = ET.SubElement(parent, q("Equipment"))
        ET.SubElement(element, q("ID")).text = item["asset_id"]
        ET.SubElement(element, q("Description")).text = item["name"]
        path = types.path(item["type"]) or [item["type"]]
        ET.SubElement(element, q("EquipmentLevel")).text = B2MML_LEVELS.get(
            path[-1], B2MML_LEVELS.get(path[0], "Other")
        )
        properties = list(item["attributes"].items())
        if item["description"]:
            properties.append(("processcraft.description", item["description"]))
        properties += [
            (
                f"{SIGNAL_PREFIX}{signal['role'] or 'signal'}",
                f"{signal['source_id']}|{signal['tag']}" if signal["source_id"] else signal["tag"],
            )
            for signal in item["signals"]
        ]
        for key, value in properties:
            prop = ET.SubElement(element, q("EquipmentProperty"))
            ET.SubElement(prop, q("ID")).text = str(key)
            holder = ET.SubElement(prop, q("Value"))
            ET.SubElement(holder, q("ValueString")).text = str(value)
        for child in children.get(item["asset_id"], []):
            add(element, child)
        ET.SubElement(element, q("EquipmentClassID")).text = item["type"]

    for item in children.get("", []):
        add(plant, item)
    return _xml(root)


def _af(platform: Any, data: dict[str, Any]) -> bytes:
    types = platform.workspace.assets.types
    root = ET.Element("AF", {"SchemaVersion": "2.0", "ExportedBy": "PlantCraft"})
    database = ET.SubElement(root, "AFDatabase")
    ET.SubElement(database, "Name").text = platform.root_asset_id
    children = _tree(data)

    def add(parent: ET.Element, item: dict[str, Any]) -> None:
        element = ET.SubElement(parent, "AFElement")
        ET.SubElement(element, "Name").text = item["asset_id"]
        ET.SubElement(element, "Description").text = item["name"]
        ET.SubElement(element, "Template").text = types.name(item["type"])
        resolved = types.resolved(item["type"]) or {"attributes": []}
        names = {entry["key"]: entry for entry in resolved["attributes"]}
        if item["description"]:
            attribute = ET.SubElement(element, "AFAttribute")
            ET.SubElement(attribute, "Name").text = "Notes"
            ET.SubElement(attribute, "Value", {"type": "String"}).text = item["description"]
        for key, value in item["attributes"].items():
            entry = names.get(key, {})
            attribute = ET.SubElement(element, "AFAttribute")
            ET.SubElement(attribute, "Name").text = entry.get("name") or key
            kind = "Double" if isinstance(value, (int, float)) else "String"
            ET.SubElement(attribute, "Value", {"type": kind}).text = str(value)
            if entry.get("unit"):
                ET.SubElement(attribute, "DefaultUOM").text = entry["unit"]
        for signal in item["signals"]:
            attribute = ET.SubElement(element, "AFAttribute")
            ET.SubElement(attribute, "Name").text = signal["role"] or signal["tag"]
            ET.SubElement(attribute, "DataReference").text = "PI Point"
            server = signal["source_id"] or "%Server%"
            ET.SubElement(attribute, "ConfigString").text = f"\\\\{server}\\{signal['tag']}"
        for child in children.get(item["asset_id"], []):
            add(element, child)

    for item in children.get("", []):
        add(database, item)
    return _xml(root)
