"""The plant map file: sources, assets, signals and loops in one file.

A map is a JSON object with one list per section, or an Excel workbook with
one sheet per section and the same columns (a CSV file holds one section,
named by the file: ``signals.csv``). Sections: ``sources``, ``assets``,
``signals`` and ``loops`` (loop assets with their signals by role,
``platform/loops.py``).

Importing is two steps. ``plan`` reads the file, checks every row and lists
what would change; ``apply`` runs that plan. Rows equal to the current state
change nothing, so importing the same file twice is a no-op, and nothing
missing from the file is removed unless ``prune`` asks for it.
"""

from __future__ import annotations

import io
import json
import math
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Protocol

from plantcraft.assets import ASSET_KINDS, KIND_LABELS
from plantcraft.i18n import t
from plantcraft.industrial_data.registry import EDITABLE, PRIMARY, SOURCE_TYPES

FORMAT = "processcraft-map"
README_SHEET = "Readme"
MAX_ROWS = 50000
SHEET_ALIASES = {"fontes": "sources", "ativos": "assets", "sinais": "signals", "malhas": "loops"}
TRUE = {"1", "true", "sim", "s", "yes", "y", "x", "verdadeiro"}
FALSE = {"0", "false", "nao", "não", "n", "no", "falso"}


@dataclass(frozen=True)
class Column:
    name: str
    description: str
    required: bool = False


@dataclass
class Change:
    action: str  # create, update, unchanged, remove
    key: str
    run: Callable[[], None] | None = None
    section: str = ""


@dataclass
class Issue:
    sheet: str
    line: int | None
    level: str  # error, warning
    message: str

    def public(self) -> dict[str, Any]:
        return {
            "sheet": self.sheet,
            "line": self.line,
            "level": self.level,
            "message": self.message,
        }


class MapSection(Protocol):
    """One section of the map; applications implement it for their own objects."""

    name: str
    label: str
    columns: list[Column]
    example: list[dict[str, Any]]

    def plan(self, rows: list[dict[str, Any]], ctx: MapContext) -> list[Change]: ...

    def stale(self, rows: list[dict[str, Any]], ctx: MapContext) -> list[Change]: ...

    def export(self, ctx: MapContext) -> list[dict[str, Any]]: ...


class MapError(ValueError):
    """The file cannot be read at all (format, size, sheets)."""


# Reading values --------------------------------------------------------------


def text(row: dict[str, Any], key: str) -> str:
    value = row.get(key)
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def flag(row: dict[str, Any], key: str, default: bool | None) -> bool | None:
    """``True``/``False`` from yes/no (or sim/não), true/false, 1/0 or x; ``None`` when unreadable."""

    value = row.get(key)
    if isinstance(value, bool):
        return value
    raw = text(row, key).lower()
    if not raw:
        return default
    if raw in TRUE:
        return True
    if raw in FALSE:
        return False
    return None


def tag_list(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    raw = "" if value is None or (isinstance(value, float) and math.isnan(value)) else str(value)
    return [
        item.strip()
        for item in raw.replace(";", "\n").replace(",", "\n").splitlines()
        if item.strip()
    ]


# The context sections share ---------------------------------------------------


@dataclass
class MapContext:
    platform: Any
    file: dict[str, list[dict[str, Any]]] = field(default_factory=dict)
    issues: list[Issue] = field(default_factory=list)
    # Who imports: named in the history of every asset record the import changes.
    actor: str = "import"
    _source_tags: dict[str, set[str] | None] = field(default_factory=dict)
    _listed: set[str] | None = None
    _claimed: dict[str, str] = field(default_factory=dict)
    _remote_sources: set[str] | None = None

    def error(self, sheet: str, row: dict[str, Any] | None, message: str) -> None:
        self.issues.append(Issue(sheet, (row or {}).get("_line"), "error", message))

    def warn(self, sheet: str, row: dict[str, Any] | None, message: str) -> None:
        self.issues.append(Issue(sheet, (row or {}).get("_line"), "warning", message))

    @property
    def root(self) -> str:
        return self.platform.root_asset_id

    def file_ids(self, section: str, key: str) -> set[str]:
        return {text(row, key) for row in self.file.get(section, []) if text(row, key)}

    def type_sheets(self) -> list[str]:
        """The sheets of the file that hold the assets of one type (``pump``, ``instrument``)."""

        types = self.platform.workspace.assets.types
        return [name for name in self.file if name not in CORE_NAMES and types.exists(name)]

    def listed_assets(self) -> set[str]:
        """Asset ids the assets sheet and the type sheets list (created before signals)."""

        if self._listed is None:
            found = self.file_ids("assets", "asset_id")
            for name in self.type_sheets():
                found |= self.file_ids(name, "asset_id")
            self._listed = found
        return self._listed

    def asset_known(self, asset_id: str, *, loops: bool = False) -> bool:
        """In the file (``loops``: also a loop of the loops sheet) or already in the tree."""

        return (
            asset_id in self.listed_assets()
            or (loops and asset_id in self.file_ids("loops", "loop_id"))
            or self.platform.workspace.assets.get(asset_id) is not None
        )

    def claim(self, sheet: str, row: dict[str, Any], asset_id: str) -> bool:
        """The first sheet that lists an asset owns it; a second listing is an error."""

        first = self._claimed.setdefault(asset_id, sheet)
        if first == sheet:
            return True
        self.error(
            sheet,
            row,
            t("Asset {asset} is already in the sheet {sheet}.", asset=asset_id, sheet=first),
        )
        return False

    def source_known(self, source_id: str) -> bool:
        if source_id in self.file_ids("sources", "source_id") or self.platform.sources.known(
            source_id
        ):
            return True
        return source_id in self.remote_sources()

    def remote_sources(self) -> set[str]:
        """With HistCraft, the sources kept there (added on its pages too): read once per file."""

        if self._remote_sources is None:
            layer = self.platform.datalayer
            found: set[str] = set()
            if layer is not None and layer.manages_sources:
                try:
                    found = {item["source_id"] for item in layer.sources()}
                except Exception:  # noqa: BLE001 - HistCraft down: only the local sources count
                    found = set()
            self._remote_sources = found
        return self._remote_sources

    def tag_source(self, tag: str) -> str:
        """Where a tag will be read from: the map's signal row, else its current route."""

        for row in self.file.get("signals", []):
            if text(row, "tag") == tag:
                return text(row, "source_id") or PRIMARY
        return self.platform.sources.source_of(tag)

    def source_tags(self, source_id: str) -> set[str] | None:
        """The tags a source serves, for warnings; ``None`` when they cannot be checked."""

        if source_id in self._source_tags:
            return self._source_tags[source_id]
        sources = self.platform.sources
        tags: set[str] | None = None
        try:
            if source_id == PRIMARY or sources.is_added(source_id):
                data = sources.data(source_id)
                if data.config.type != "pi_web_api":  # PI is not browsed: it lists its own tags
                    tags = {item["tag"] for item in data.signals()}
                    if source_id == PRIMARY:
                        tags |= set(data.config.mapping.values())
            else:
                # A new CSV source in the same file is quick to open; servers are checked after import.
                row = next(
                    (
                        item
                        for item in self.file.get("sources", [])
                        if text(item, "source_id") == source_id
                    ),
                    None,
                )
                if row is not None and text(row, "type") == "csv" and text(row, "path"):
                    from plantcraft.industrial_data.sources import build_source

                    config = sources.validate("csv", {"path": text(row, "path")})
                    tags = {item.tag for item in build_source(config).list_signals()}
                elif row is not None:
                    self.warn(
                        "signals",
                        None,
                        t(
                            "The tags of the new source {source} were not checked; "
                            "after importing, a preview of the same file checks them all.",
                            source=source_id,
                        ),
                    )
        except Exception as exc:  # the source may be offline: say so once
            self.warn(
                "signals",
                None,
                t(
                    "Could not check the tags of source {source}: {error}",
                    source=source_id,
                    error=exc,
                ),
            )
        self._source_tags[source_id] = tags
        return tags

    def check_tag(
        self, sheet: str, row: dict[str, Any], tag: str, source_id: str | None = None
    ) -> None:
        source_id = source_id or self.tag_source(tag)
        known = self.source_tags(source_id)
        if known is not None and tag not in known:
            self.warn(
                sheet,
                row,
                t("Tag {tag} was not found in source {source}.", tag=tag, source=source_id),
            )


# Core sections ----------------------------------------------------------------


def _columns(spec: tuple[tuple[str, str, bool], ...]) -> list[Column]:
    """Columns with their descriptions in the current language."""

    return [Column(name, t(description), required) for name, description, required in spec]


class SourcesSection:
    name = "sources"
    # English texts, translated each time the section is read (a request's language).
    COLUMNS = (
        (
            "source_id",
            "Source id, used in the signals sheet. 'scenario' is the scenario file's source (read-only).",
            True,
        ),
        ("name", "Name shown in the platform.", True),
        ("type", "opcua, pi_web_api or csv.", True),
        ("endpoint_url", "OPC UA: server address (opc.tcp://host:4840/path).", False),
        ("namespace_uri", "OPC UA: namespace of the NodeIds (urn:...).", False),
        ("base_url", "PI Web API: base address (https://server/piwebapi).", False),
        ("server", "PI Web API: name of the PI Data Archive.", False),
        ("path", "CSV: path of the file on the server.", False),
        ("auth_type", "OPC UA: anonymous or basic. PI: kerberos, ntlm or basic.", False),
        ("username", "User, when authentication asks for one.", False),
        (
            "password_env",
            "Server environment variable holding the password. Never write the password in the file; it can also be typed in the Historian.",
            False,
        ),
        ("timeout_seconds", "Maximum wait for each read, in seconds.", False),
        (
            "tags",
            "PI Web API: the source's tags separated by semicolons (PI is not browsed).",
            False,
        ),
    )
    example = [
        {
            "source_id": "opc-utilities",
            "name": "OPC UA Utilities",
            "type": "opcua",
            "endpoint_url": "opc.tcp://10.0.0.15:4840/utilities/",
            "namespace_uri": "urn:utilities:opc",
            "auth_type": "anonymous",
            "timeout_seconds": 30,
        }
    ]

    @property
    def label(self) -> str:
        return t("Sources")

    @property
    def columns(self) -> list[Column]:
        return _columns(self.COLUMNS)

    def _settings(self, row: dict[str, Any], type: str) -> dict[str, Any]:
        settings: dict[str, Any] = {}
        for key in EDITABLE.get(type, ()):
            if key == "tags":
                if tag_list(row.get("tags")):
                    settings["tags"] = tag_list(row.get("tags"))
            elif key == "verify_ssl":
                value = flag(row, key, None)
                if value is not None:
                    settings[key] = value
            elif text(row, key):
                settings[key] = text(row, key)
        return settings

    def plan(self, rows: list[dict[str, Any]], ctx: MapContext) -> list[Change]:
        registry = ctx.platform.sources
        changes: list[Change] = []
        seen: set[str] = set()
        for row in rows:
            source_id, name, type = text(row, "source_id"), text(row, "name"), text(row, "type")
            if not source_id:
                ctx.error(self.name, row, t("Enter source_id."))
                continue
            if source_id in seen:
                ctx.error(
                    self.name, row, t("Source {source} appears more than once.", source=source_id)
                )
                continue
            seen.add(source_id)
            if source_id == PRIMARY:
                changes.append(Change("unchanged", source_id))
                continue
            if type not in SOURCE_TYPES:
                ctx.error(
                    self.name,
                    row,
                    t("Invalid source type: '{type}'. Use opcua, pi_web_api or csv.", type=type),
                )
                continue
            if not name:
                ctx.error(self.name, row, t("Enter the source name."))
                continue
            settings = self._settings(row, type)
            try:
                config = registry.validate(type, settings)
            except ValueError as exc:
                ctx.error(self.name, row, str(exc))
                continue
            stored = config.model_dump(
                mode="json", include=set(EDITABLE[type]) | {"timezone", "history_cache_hours"}
            )
            current = registry.store.source(source_id)
            needs_password = config.auth_type in {"basic", "ntlm"} and not config.password_env
            if needs_password and not (current and current.get("secret")):
                ctx.warn(
                    self.name,
                    row,
                    t(
                        "Authentication needs a password: type it in the Historian's source after importing."
                    ),
                )
            if (
                current is not None
                and current["name"] == name
                and current["type"] == type
                and current["settings"] == stored
            ):
                changes.append(Change("unchanged", source_id))
                continue

            def run(
                source_id: str = source_id,
                name: str = name,
                type: str = type,
                settings: dict = settings,
            ) -> None:
                layer = ctx.platform.datalayer
                if layer is not None and layer.manages_sources:
                    # With HistCraft, sources live there: the import creates or edits them there.
                    layer.save_source(
                        source_id,
                        name=name,
                        type=type,
                        settings=registry.validate(type, settings).model_dump(
                            mode="json",
                            include=set(EDITABLE[type]) | {"timezone", "history_cache_hours"},
                        ),
                        password=None,
                        create=layer.source(source_id) is None,
                    )
                    return
                registry.save(
                    source_id,
                    name=name,
                    type=type,
                    settings=settings,
                    password=None,
                    actor="import",
                )

            changes.append(Change("update" if current else "create", source_id, run))
        return changes

    def stale(self, rows: list[dict[str, Any]], ctx: MapContext) -> list[Change]:
        listed = {text(row, "source_id") for row in rows}
        registry = ctx.platform.sources
        return [
            Change(
                "remove",
                item["source_id"],
                lambda source_id=item["source_id"]: registry.remove(source_id),
            )
            for item in registry.store.sources()
            if item["source_id"] not in listed
        ]

    def export(self, ctx: MapContext) -> list[dict[str, Any]]:
        rows = []
        for item in ctx.platform.sources.list():
            settings = dict(item["settings"])
            if isinstance(settings.get("tags"), list):
                settings["tags"] = "; ".join(settings["tags"])
            elif "tags" in settings:
                settings.pop("tags")
            row = {"source_id": item["source_id"], "name": item["name"], "type": item["type"]}
            row.update({column.name: settings.get(column.name, "") for column in self.columns[3:]})
            rows.append(row)
        return rows


class AssetsSection:
    name = "assets"
    COLUMNS = (
        ("asset_id", "Asset id (no spaces: U-200, B-201A).", True),
        ("name", "Name shown in the asset tree.", True),
        (
            "kind",
            "Asset type: unit, area, equipment, loop or another type registered in the platform.",
            True,
        ),
        ("parent_id", "Parent asset. Empty: right under the site.", False),
        ("description", "Free description.", False),
        (
            "attributes",
            'JSON: object with the asset\'s attribute values by key ({"manufacturer": "ACME"}). In the spreadsheet, each extra column of the sheet is an attribute. The sheet of each type has one column per attribute.',
            False,
        ),
    )

    @property
    def label(self) -> str:
        return t("Assets")

    @property
    def columns(self) -> list[Column]:
        return _columns(self.COLUMNS)

    # The template puts each asset in the sheet of its type; this sheet takes any type.
    example: list[dict[str, Any]] = []

    def plan(self, rows: list[dict[str, Any]], ctx: MapContext) -> list[Change]:
        store = ctx.platform.workspace.assets
        valid: dict[str, dict[str, Any]] = {}
        values_of: dict[str, dict[str, Any]] = {}
        ignored: set[str] = set()
        for row in rows:
            asset_id, name, kind = (
                text(row, "asset_id"),
                text(row, "name"),
                text(row, "kind").lower(),
            )
            if not asset_id or not name:
                ctx.error(self.name, row, t("Enter asset_id and name."))
                continue
            if any(char.isspace() for char in asset_id):
                ctx.error(
                    self.name, row, t("The asset_id '{asset}' cannot have spaces.", asset=asset_id)
                )
                continue
            if asset_id in valid:
                ctx.error(
                    self.name, row, t("Asset {asset} appears more than once.", asset=asset_id)
                )
                continue
            if not _is_kind(store, kind) or kind == "site":
                ctx.error(
                    self.name,
                    row,
                    t(
                        "Invalid asset type: '{kind}'. Use {kinds}.",
                        kind=kind,
                        kinds=_kind_names(store),
                    ),
                )
                continue
            if asset_id == ctx.root or not ctx.claim(self.name, row, asset_id):
                continue
            values, unknown, problems = record_values(
                store, kind, _attributes(row, self.columns), ctx
            )
            ignored.update(unknown)
            if problems:
                for problem in problems:
                    ctx.error(self.name, row, problem)
                continue
            valid[asset_id] = row
            values_of[asset_id] = values
        if ignored:
            ctx.warn(
                self.name,
                None,
                t(
                    "Not attributes of the asset's type, ignored: {attributes}. Add them to the "
                    "type in Assets › Types first.",
                    attributes=", ".join(sorted(ignored)),
                ),
            )
        # Parents first, so each asset is created under an existing one.
        ordered: list[str] = []
        placing: set[str] = set()

        def place(asset_id: str) -> None:
            if asset_id in ordered:
                return
            if asset_id in placing:
                ctx.error(
                    self.name,
                    valid[asset_id],
                    t("Asset {asset} ends up inside itself (cycle in parent_id).", asset=asset_id),
                )
                return
            placing.add(asset_id)
            parent = text(valid[asset_id], "parent_id")
            if parent in valid:
                place(parent)
            placing.discard(asset_id)
            if asset_id not in ordered:
                ordered.append(asset_id)

        for asset_id in valid:
            place(asset_id)
        changes: list[Change] = []
        for asset_id in ordered:
            row = valid[asset_id]
            parent = text(row, "parent_id") or ctx.root
            if parent not in valid and parent != ctx.root and store.get(parent) is None:
                ctx.warn(
                    self.name,
                    row,
                    t(
                        "Parent asset {parent} does not exist; {asset} goes under the site.",
                        parent=parent,
                        asset=asset_id,
                    ),
                )
                parent = ctx.root
            wanted = {
                "name": text(row, "name"),
                "kind": text(row, "kind").lower(),
                "parent_id": parent,
                "description": text(row, "description"),
            }
            current = store.get(asset_id)
            record = changed_values(store, asset_id, values_of[asset_id])
            if (
                current is not None
                and all(current.get(key) == value for key, value in wanted.items())
                and not record
            ):
                changes.append(Change("unchanged", asset_id))
                continue

            def run(
                asset_id: str = asset_id,
                wanted: dict = wanted,
                current: dict | None = current,
                record: dict = record,
            ) -> None:
                store.save(
                    asset_id, origin="manual", position=(current or {}).get("position", 0), **wanted
                )
                write_record(ctx, asset_id, record)

            changes.append(Change("update" if current else "create", asset_id, run))
        return changes

    def stale(self, rows: list[dict[str, Any]], ctx: MapContext) -> list[Change]:
        # Assets of a type with its own sheet in the file are that sheet's to remove.
        listed = ctx.listed_assets()
        own_sheets = set(ctx.type_sheets())
        store = ctx.platform.workspace.assets
        return [
            Change(
                "remove", item["asset_id"], lambda asset_id=item["asset_id"]: store.delete(asset_id)
            )
            for item in store.list()
            if item["origin"] == "manual"
            and item["asset_id"] != ctx.root
            and item["kind"] not in own_sheets
            and item["asset_id"] not in listed
        ]

    def export(self, ctx: MapContext) -> list[dict[str, Any]]:
        # Every asset is exported in the sheet of its type (``typesheets.py``).
        return []


class SignalsSection:
    name = "signals"
    COLUMNS = (
        (
            "tag",
            "Tag in the source (text NodeId in OPC UA, point name in PI, column in CSV).",
            True,
        ),
        ("source_id", "Source that serves the tag. Empty: the scenario's source.", False),
        ("name", "Readable name of the signal.", False),
        ("unit", "Engineering unit (m³/h, °C, %).", False),
        ("asset_id", "Asset the signal belongs to (assets sheet or an existing asset).", False),
        (
            "role",
            "The signal's role on the asset: pv, sp, op, mode, pressure, vibration... (free text).",
            False,
        ),
        ("monitor", "yes/no: the platform follows the signal (data quality). Empty: yes.", False),
    )

    @property
    def label(self) -> str:
        return t("Signals")

    @property
    def columns(self) -> list[Column]:
        return _columns(self.COLUMNS)

    @property
    def example(self) -> list[dict[str, Any]]:
        return [
            {
                "tag": "UT.B201A.Flow",
                "source_id": "opc-utilities",
                "name": t("Discharge flow"),
                "unit": "m³/h",
                "asset_id": "B-201A",
                "role": "flow",
                "monitor": True,
            },
            {
                "tag": "UT.FIC2001.PV",
                "source_id": "opc-utilities",
                "name": "FIC-2001 PV",
                "unit": "m³/h",
                "asset_id": "",
                "role": "",
                "monitor": True,
            },
        ]

    def plan(self, rows: list[dict[str, Any]], ctx: MapContext) -> list[Change]:
        registry = ctx.platform.sources
        store = ctx.platform.workspace.assets
        configured = set(ctx.platform.scenario.source.tags)
        routes = registry.routes()
        tags = sorted({text(row, "tag") for row in rows if text(row, "tag")})
        current_bindings = store.signal_assets(tags)
        try:
            primary_info = (
                {item["tag"]: item for item in registry.primary_data.signals()} if tags else {}
            )
        except Exception:
            primary_info = {}
        changes: list[Change] = []
        source_of: dict[str, str] = {}
        bound: set[tuple[str, str]] = set()
        layer = ctx.platform.datalayer
        remote = layer is not None and layer.manages_sources
        # With HistCraft, the source's signals are collected there: one call per source.
        collect: dict[tuple[str, bool], list[str]] = {}
        for row in rows:
            tag, source_id = text(row, "tag"), text(row, "source_id") or PRIMARY
            asset_id, role = text(row, "asset_id"), text(row, "role")
            monitor = flag(row, "monitor", True)
            if not tag:
                ctx.error(self.name, row, t("Enter the tag."))
                continue
            if monitor is None:
                ctx.error(
                    self.name,
                    row,
                    t(
                        "Invalid value in monitor: '{value}'. Use yes or no.",
                        value=text(row, "monitor"),
                    ),
                )
                continue
            if not ctx.source_known(source_id):
                ctx.error(
                    self.name,
                    row,
                    t(
                        "Source {source} does not exist and is not in the sources sheet.",
                        source=source_id,
                    ),
                )
                continue
            if source_of.setdefault(tag, source_id) != source_id:
                ctx.error(
                    self.name,
                    row,
                    t(
                        "Tag {tag} appears with two sources ({first} and {second}).",
                        tag=tag,
                        first=source_of[tag],
                        second=source_id,
                    ),
                )
                continue
            if asset_id and not ctx.asset_known(asset_id):
                ctx.error(
                    self.name,
                    row,
                    t(
                        "Asset {asset} does not exist and is not in the assets sheet.",
                        asset=asset_id,
                    ),
                )
                continue
            if asset_id and (asset_id, tag) in bound:
                ctx.error(
                    self.name,
                    row,
                    t("Tag {tag} appears twice on asset {asset}.", tag=tag, asset=asset_id),
                )
                continue
            if asset_id:
                bound.add((asset_id, tag))
            ctx.check_tag(self.name, row, tag, source_id)
            if source_id == PRIMARY and tag in configured and not monitor:
                ctx.warn(
                    self.name,
                    row,
                    t("Tag {tag} is in the scenario file and stays followed.", tag=tag),
                )
                monitor = True
            name, unit = text(row, "name"), text(row, "unit")
            route = routes.get(tag)
            info = route or primary_info.get(tag) or {}
            current_source = route["source_id"] if route else PRIMARY
            monitored = bool(route and route["monitored"]) or (
                current_source == PRIMARY and tag in configured
            )
            same_route = (
                current_source == source_id
                and monitored == monitor
                and (not name or name == (info.get("name") or ""))
                and (not unit or unit == (info.get("unit") or ""))
            )
            if source_id == PRIMARY and not route and not monitor and not name and not unit:
                same_route = True  # nothing to remember about a scenario tag nobody monitors
            uses = current_bindings.get(tag, [])
            same_binding = not asset_id or any(
                item["asset_id"] == asset_id and item["role"] == role for item in uses
            )
            if same_route and same_binding:
                changes.append(Change("unchanged", tag))
                continue

            def run(
                tag: str = tag,
                source_id: str = source_id,
                name: str = name,
                unit: str = unit,
                monitor: bool = monitor,
                asset_id: str = asset_id,
                role: str = role,
                same_route: bool = same_route,
            ) -> None:
                if not same_route and remote:
                    collect.setdefault((source_id, monitor), []).append(tag)
                elif not same_route:
                    registry.use(
                        source_id,
                        [{"tag": tag, "name": name or None, "unit": unit or None}],
                        monitored=monitor,
                    )
                if asset_id:
                    store.bind_signals(
                        asset_id,
                        [
                            {
                                "signal_id": tag,
                                "role": role,
                                "name": name,
                                "unit": unit,
                                "source_id": source_id,
                            }
                        ],
                    )

            changes.append(Change("update" if route or uses else "create", tag, run))
        runs = [change for change in changes if change.run is not None]
        if remote and runs:
            last, inner = runs[-1], runs[-1].run

            def finish(inner: Callable[[], None] = inner) -> None:
                inner()
                for (source_id, monitor), batch in collect.items():
                    if monitor:
                        layer.add_signals(source_id, batch)
                    else:
                        layer.stop_signals(source_id, batch)

            last.run = finish
        return changes

    def stale(self, rows: list[dict[str, Any]], ctx: MapContext) -> list[Change]:
        listed = {text(row, "tag") for row in rows}
        pairs = {(text(row, "asset_id"), text(row, "tag")) for row in rows}
        registry = ctx.platform.sources
        store = ctx.platform.workspace.assets
        changes = [
            Change("remove", tag, lambda tag=tag: registry.forget([tag]))
            for tag in registry.routes()
            if tag not in listed
        ]
        changes.extend(
            Change(
                "remove",
                f"{item['asset_id']} · {item['signal_id']}",
                lambda pair=(item["asset_id"], item["signal_id"]): store.unbind_signals([pair]),
            )
            for item in store.bindings(origin="manual")
            if (item["asset_id"], item["signal_id"]) not in pairs
        )
        return changes

    def export(self, ctx: MapContext) -> list[dict[str, Any]]:
        platform = ctx.platform
        registry = platform.sources
        routes = registry.routes()
        try:
            info = {item["tag"]: item for item in registry.primary_data.signals()}
        except Exception:
            info = {}
        bindings: dict[str, list[dict[str, Any]]] = {}
        for item in platform.workspace.assets.bindings(origin="manual"):
            bindings.setdefault(item["signal_id"], []).append(item)
        tags = list(dict.fromkeys([*platform.scenario.source.tags, *routes, *bindings]))
        rows = []
        for tag in tags:
            route = routes.get(tag)
            meta = route or info.get(tag) or {}
            base = {
                "tag": tag,
                "source_id": route["source_id"] if route else PRIMARY,
                "name": meta.get("name") or "",
                "unit": meta.get("unit") or "",
                "monitor": bool(route["monitored"])
                if route
                else tag in platform.scenario.source.tags,
            }
            for binding in bindings.get(tag) or [None]:
                rows.append(
                    {
                        **base,
                        "asset_id": binding["asset_id"] if binding else "",
                        "role": binding["role"] if binding else "",
                    }
                )
        return [{column.name: row.get(column.name, "") for column in self.columns} for row in rows]


def record_values(
    store: Any, kind: str, given: dict[str, Any], ctx: MapContext
) -> tuple[dict[str, Any], list[str], list[str]]:
    """Attribute values of a row checked against the type: (values, unknown keys, problems)."""

    from plantcraft.assets.types import AssetTypeError, check_value

    resolved = store.types.resolved(kind) or {"attributes": []}
    by_key = {entry["key"]: entry for entry in resolved["attributes"]}
    values: dict[str, Any] = {}
    unknown: list[str] = []
    problems: list[str] = []
    for key, raw in given.items():
        entry = by_key.get(str(key).strip().lower())
        if entry is None:
            unknown.append(str(key)[:40])
            continue
        try:
            value = check_value(entry, raw)
        except AssetTypeError as exc:
            problems.append(str(exc))
            continue
        if value is None:
            continue  # an empty cell leaves the value as it is
        if entry["data_type"] == "asset" and not ctx.asset_known(value, loops=True):
            problems.append(
                t("{key}: asset {asset} not found.", key=entry["key"], asset=str(value)[:120])
            )
            continue
        values[entry["key"]] = value
    return values, unknown, problems


def changed_values(store: Any, asset_id: str, values: dict[str, Any]) -> dict[str, Any]:
    stored = store.stored_attributes([asset_id]).get(asset_id, {})
    return {key: value for key, value in values.items() if stored.get(key) != value}


def write_record(ctx: MapContext, asset_id: str, values: dict[str, Any]) -> None:
    """Store checked values and put the change in the asset's history."""

    if not values:
        return
    platform = ctx.platform
    changes = platform.workspace.assets.set_attributes(
        asset_id, values, actor=ctx.actor, check_refs=False
    )
    platform.record_attribute_changes(asset_id, ctx.actor, changes)


def _is_kind(store: Any, kind: str) -> bool:
    checker = getattr(store, "is_kind", None)
    return bool(kind) and (checker(kind) if checker else kind in ASSET_KINDS)


def _kind_names(store: Any) -> str:
    kinds = getattr(store, "kinds", None)
    if kinds is not None:
        names = [item if isinstance(item, str) else str(item.get("kind")) for item in kinds()]
        return ", ".join(name for name in names if name != "site")
    return ", ".join(f"{item} ({t(KIND_LABELS[item]).lower()})" for item in ASSET_KINDS[1:])


def _attributes(row: dict[str, Any], columns: list[Column]) -> dict[str, Any]:
    """An ``attributes`` object (JSON) or the sheet's extra columns (Excel, CSV)."""

    known = {column.name for column in columns} | {"_line"}
    found = dict(row["attributes"]) if isinstance(row.get("attributes"), dict) else {}
    if isinstance(row.get("attributes"), str) and row["attributes"].strip():
        try:
            parsed = json.loads(row["attributes"])
            found.update(parsed if isinstance(parsed, dict) else {})
        except json.JSONDecodeError:
            found["attributes"] = row["attributes"]
    found.update({key: value for key, value in row.items() if key not in known})
    return found


def _tree_order(items: list[dict[str, Any]], root: str) -> list[dict[str, Any]]:
    children: dict[str | None, list[dict[str, Any]]] = {}
    ids = {item["asset_id"] for item in items}
    for item in items:
        parent = item["parent_id"] if item["parent_id"] in ids else None
        children.setdefault(parent, []).append(item)
    ordered: list[dict[str, Any]] = []
    pending = list(reversed(children.get(None, [])))
    while pending:
        item = pending.pop()
        ordered.append(item)
        pending.extend(reversed(children.get(item["asset_id"], [])))
    return ordered


CORE_SECTIONS: tuple[MapSection, ...] = (SourcesSection(), AssetsSection(), SignalsSection())
# Sheets that are not the sheet of a type.
CORE_NAMES = {"sources", "assets", "signals", "loops", "relations"}


def sections(platform: Any) -> list[MapSection]:
    """Sources, assets, one sheet per type, signals, loops and relations, in apply order.

    Assets come before signals (a signal binds to an existing asset) and relations
    last (they link assets of every sheet).
    """

    # Loops and type sheets build on this module's helpers, so they are imported here.
    from plantcraft.platform.loops import LoopsSection
    from plantcraft.platform.typesheets import RelationsSection, type_sections

    sources, assets, signals = CORE_SECTIONS
    return [
        sources,
        assets,
        *type_sections(platform),
        signals,
        LoopsSection(platform),
        RelationsSection(),
    ]


# Files -------------------------------------------------------------------------


def read_file(content: bytes, filename: str, names: list[str]) -> dict[str, list[dict[str, Any]]]:
    """Rows per section from a .json, .xlsx or .csv file; each row knows its line."""

    suffix = Path(filename or "").suffix.lower()
    if suffix == ".json" or (not suffix and content.lstrip().startswith(b"{")):
        return _read_json(content, names)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_excel(content, names)
    if suffix == ".csv":
        section = SHEET_ALIASES.get(Path(filename).stem.lower(), Path(filename).stem.lower())
        if section not in names:
            raise MapError(
                t(
                    "Name the CSV after the sheet it holds: {names}.",
                    names=", ".join(f"{name}.csv" for name in names),
                )
            )
        import pandas as pd

        try:
            frame = pd.read_csv(
                io.BytesIO(content), dtype=str, keep_default_na=False, sep=None, engine="python"
            )
        except Exception as exc:
            raise MapError(t("Could not read the CSV: {error}", error=exc)) from exc
        return {section: _frame_rows(frame)}
    raise MapError(t("Send a .json, .xlsx or .csv file."))


def _read_json(content: bytes, names: list[str]) -> dict[str, list[dict[str, Any]]]:
    try:
        payload = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise MapError(t("Invalid JSON: {error}", error=exc)) from exc
    if not isinstance(payload, dict):
        raise MapError(
            t(
                "The JSON must be an object with the lists sources, assets, signals and loops "
                "(and one list per asset type, and relations)."
            )
        )
    found: dict[str, list[dict[str, Any]]] = {}
    for key, value in payload.items():
        section = SHEET_ALIASES.get(key.lower(), key.lower())
        if section not in names:
            continue
        if not isinstance(value, list) or not all(isinstance(item, dict) for item in value):
            raise MapError(t("'{key}' must be a list of objects.", key=key))
        found[section] = [{**item, "_line": index + 1} for index, item in enumerate(value)]
    return found


def _read_excel(content: bytes, names: list[str]) -> dict[str, list[dict[str, Any]]]:
    import pandas as pd

    _require_openpyxl()
    try:
        sheets = pd.read_excel(
            io.BytesIO(content), sheet_name=None, dtype=object, engine="openpyxl"
        )
    except Exception as exc:
        raise MapError(t("Could not read the spreadsheet: {error}", error=exc)) from exc
    found: dict[str, list[dict[str, Any]]] = {}
    for sheet, frame in sheets.items():
        section = SHEET_ALIASES.get(str(sheet).strip().lower(), str(sheet).strip().lower())
        if section in names:
            found[section] = _frame_rows(frame)
    if not found:
        raise MapError(
            t("The spreadsheet has none of the sheets {sheets}.", sheets=", ".join(names))
        )
    return found


def _frame_rows(frame: Any) -> list[dict[str, Any]]:
    frame.columns = [str(column).strip().lower() for column in frame.columns]
    rows = []
    for index, record in enumerate(frame.to_dict(orient="records")):
        clean = {key: value for key, value in record.items() if not _blank(value)}
        if clean:
            # Line 1 is the header row.
            rows.append({**clean, "_line": index + 2})
    if len(rows) > MAX_ROWS:
        raise MapError(t("Each sheet takes up to {count} rows.", count=MAX_ROWS))
    return rows


def _blank(value: Any) -> bool:
    return (
        value is None
        or (isinstance(value, float) and math.isnan(value))
        or (isinstance(value, str) and not value.strip())
    )


def _require_openpyxl() -> None:
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise MapError(
            t(
                "Reading and writing Excel spreadsheets needs the openpyxl package: python3 -m pip install -e '.[server]'."
            )
        ) from exc


# Plan, apply, export -------------------------------------------------------------


def plan(
    platform: Any, content: bytes, filename: str, *, prune: bool = False, actor: str = "import"
) -> tuple[dict[str, Any], list[Change]]:
    """What importing the file would change, and the changes to run."""

    all_sections = sections(platform)
    names = [section.name for section in all_sections]
    rows = read_file(content, filename, names)
    ctx = MapContext(platform, rows, actor=actor)
    if not any(rows.values()):
        ctx.error(
            "file", None, t("The file has no rows in the sheets {sheets}.", sheets=", ".join(names))
        )
    summary: list[dict[str, Any]] = []
    ordered: list[Change] = []
    removals: list[Change] = []
    for section in all_sections:
        section_rows = rows.get(section.name)
        if section_rows is None:
            continue
        changes = section.plan(section_rows, ctx)
        stale = section.stale(section_rows, ctx) if prune else []
        for change in [*changes, *stale]:
            change.section = section.name
        ordered.extend(changes)
        removals = stale + removals  # objects go away in the reverse order they were created
        counts = {
            action: sum(1 for item in changes if item.action == action)
            for action in ("create", "update", "unchanged")
        }
        summary.append(
            {
                "name": section.name,
                "label": section.label,
                "rows": len(section_rows),
                **counts,
                "remove": len(stale),
            }
        )
    errors = [item for item in ctx.issues if item.level == "error"]
    report = {
        "file": filename,
        "sections": summary,
        "issues": [item.public() for item in ctx.issues],
        "errors": len(errors),
        "warnings": len(ctx.issues) - len(errors),
        "changes": [
            {"section": item.section, "action": item.action, "key": item.key}
            for item in [*ordered, *removals]
            if item.action != "unchanged"
        ][:500],
        "can_apply": not errors,
        "prune": prune,
    }
    return report, [*ordered, *removals]


def apply(
    platform: Any, content: bytes, filename: str, *, prune: bool = False, actor: str = "import"
) -> dict[str, Any]:
    report, changes = plan(platform, content, filename, prune=prune, actor=actor)
    if not report["can_apply"]:
        return {**report, "applied": False}
    for change in changes:
        if change.run is not None:
            change.run()
    return {**report, "applied": True}


def export(platform: Any) -> dict[str, Any]:
    """The current configuration in the map format (loops first, so assets can leave them out)."""

    ctx = MapContext(platform)
    payload: dict[str, Any] = {"format": FORMAT}
    found: dict[str, list[dict[str, Any]]] = {}
    for section in reversed(sections(platform)):
        found[section.name] = section.export(ctx)
        ctx.file[section.name] = found[section.name]
    for section in sections(platform):
        payload[section.name] = found[section.name]
    return payload


def template(platform: Any) -> dict[str, Any]:
    payload: dict[str, Any] = {"format": FORMAT}
    for section in sections(platform):
        payload[section.name] = [dict(row) for row in section.example]
    return payload


def workbook(platform: Any, payload: dict[str, Any]) -> bytes:
    """An Excel file with one sheet per section and a Readme sheet that explains every column."""

    _require_openpyxl()
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    readme = book.active
    readme.title = t(README_SHEET)
    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DCEFF0")
    readme.append([t("Plant map · PlantCraft")])
    readme["A1"].font = Font(bold=True, size=14)
    for line in (
        "One sheet per list: sources, assets, one sheet per asset type (pump, instrument...), signals, loops and relations.",
        "A type's sheet has one column per attribute and one per signal role (the signal's tag). An empty cell leaves the value as it is.",
        "The first row of each sheet holds the column names; do not translate them. Empty rows are ignored.",
        "Import in Settings › Configuration file. The preview shows what will be created or changed and the errors by sheet and row.",
        "Importing the same file again changes nothing. What is not in the file is only removed if you tick that option.",
        "Never write passwords here: use password_env (environment variable on the server) or type the password in the Historian.",
    ):
        readme.append([t(line)])
    readme.append([])
    readme.append([t("Sheet"), t("Column"), t("Required"), t("What it is")])
    for cell in readme[readme.max_row]:
        cell.font = bold
        cell.fill = header_fill
    for section in sections(platform):
        for column in section.columns:
            readme.append(
                [section.name, column.name, t("yes") if column.required else "", column.description]
            )
    readme.column_dimensions["A"].width = 12
    readme.column_dimensions["B"].width = 18
    readme.column_dimensions["C"].width = 12
    readme.column_dimensions["D"].width = 110
    for section in sections(platform):
        sheet = book.create_sheet(section.name)
        names = [column.name for column in section.columns]
        sheet.append(names)
        for cell in sheet[1]:
            cell.font = bold
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for row in payload.get(section.name) or []:
            sheet.append([_cell(row.get(name)) for name in names])
        for index, name in enumerate(names, start=1):
            width = max(
                [
                    len(name),
                    *(
                        len(str(_cell(row.get(name)) or ""))
                        for row in payload.get(section.name) or []
                    ),
                ]
            )
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(
                max(width + 2, 10), 60
            )
        sheet.freeze_panes = "A2"
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def _cell(value: Any) -> Any:
    if isinstance(value, bool):
        return t("yes") if value else t("no")
    if isinstance(value, list):
        return "; ".join(str(item) for item in value)
    return value


def format_reference(platform: Any) -> list[dict[str, Any]]:
    """Every section and column, for the import page and the help."""

    return [
        {
            "name": section.name,
            "label": section.label,
            "columns": [
                {
                    "name": column.name,
                    "required": column.required,
                    "description": column.description,
                }
                for column in section.columns
            ],
        }
        for section in sections(platform)
    ]
