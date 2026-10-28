"""Shared fixtures: the fake PlantCraft of ``tests/plant.py``, a PlantCraft Documents wired to it,
and small documents of the test plant (``FILES``: a PDF incident report, text and Word made here)."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from plantcraft_sdk.testing import BASE_URL

from doccraft.api.app import create_app
from doccraft.api.runtime import Runtime
from doccraft.config import Settings
from tests.plant import APP_KEY, make_fake

FILES = Path(__file__).parent / "files"


def make_runtime(tmp_path, fake, *, app_key: bool = True, **settings) -> Runtime:
    """No indexing thread: a test reads the queue itself (``index``), so nothing races."""

    return Runtime(
        Settings(processcraft_url=BASE_URL, data_dir=tmp_path, embedding_dir="", **settings),
        verifier=fake.verifier("doccraft"),
        service_client=fake.client(None, app="doccraft", key=APP_KEY) if app_key else None,
        http_client=fake.http_client(),
        background=False,
    )


@pytest.fixture
def fake():
    return make_fake()


@pytest.fixture
def runtime(tmp_path, fake):
    rt = make_runtime(tmp_path, fake)
    yield rt
    rt.shutdown()


def client_for(runtime: Runtime, fake, username: str | None = "engineer") -> TestClient:
    client = TestClient(
        create_app(runtime), base_url="http://testserver/doc/", raise_server_exceptions=False
    )
    if username:
        client.headers["Authorization"] = f"Bearer {fake.token(username)}"
    return client


@pytest.fixture
def client(runtime, fake):
    with client_for(runtime, fake) as test_client:
        yield test_client


GUIDE = """# Control loop troubleshooting guide

## 4.1 Valve stiction

A loop that cycles with a saw-tooth controller output and a square-wave flow has a sticking valve.
FIC-1001 cycles this way when the packing of its valve is too tight. Put the loop in manual and
move the output in steps of 0.5 % to measure the dead band; a dead band above 1 % needs maintenance.
Check the positioner before retuning: a slower integral action only hides the stiction.

## 4.2 Sensor drift

TIC-1209 drifts slowly when its thermowell fouls. Compare it with the spare element every month and
recalibrate when the two differ by more than 2 degrees. Retuning does not fix a drifting sensor.
"""


def send(client: TestClient, filename: str, data: bytes, **query) -> dict:
    """Send a file the way the page does (the body is the file); the answer's document."""

    answer = client.post("/api/documents", params={"filename": filename, **query}, content=data)
    assert answer.status_code == 201, answer.text
    return answer.json()


def index(runtime: Runtime) -> int:
    """Read every queued file now, as the indexing thread would."""

    return runtime.library.work_once()


def guide(client: TestClient, runtime: Runtime, **query) -> dict:
    doc = send(
        client,
        "troubleshooting.md",
        GUIDE.encode(),
        number="DOC-TG-001",
        kind="Troubleshooting guide",
        **query,
    )
    index(runtime)
    return doc


def docx_bytes(heading: str, *paragraphs: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(heading, level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def pytest_sessionfinish(session, exitstatus):
    """The app's token used only what app.toml [permissions] declares (PlantCraft answers 403)."""

    from pathlib import Path

    from plantcraft_sdk import apptest

    problems = apptest.permission_problems(Path(__file__).resolve().parents[1], "doccraft")
    if problems and exitstatus == 0:
        print("\n" + "\n".join(problems))
        session.exitstatus = 1
